import sys
from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1]
maxrows = int(sys.argv[2]) if len(sys.argv) > 2 else 4
d = Document(path)
body = d.element.body
ti = 0
for el in body.iterchildren():
    tag = el.tag.split('}')[1]
    if tag == 'p':
        style = el.find(qn('w:pPr'))
        sname = ''
        if style is not None and style.find(qn('w:pStyle')) is not None:
            sname = style.find(qn('w:pStyle')).get(qn('w:val'))
        numpr = ' [NUM]' if el.find('.//' + qn('w:numPr')) is not None else ''
        text = ''.join(t.text or '' for t in el.iter(qn('w:t')))
        if text.strip() or numpr:
            print(f'P[{sname}]{numpr}: {text[:160]}')
    elif tag == 'tbl':
        t = d.tables[ti]; ti += 1
        print(f'--- TABLE {ti} rows={len(t.rows)} cols={len(t.columns)}')
        for ri, r in enumerate(t.rows[:maxrows]):
            cells = []
            for c in r.cells:
                parts = []
                for p in c.paragraphs:
                    s = p.style.name if p.style is not None else ''
                    n = '•' if p._p.find('.//' + qn('w:numPr')) is not None else ''
                    parts.append(f'{n}<{s}>{p.text[:90]}')
                cells.append(' | '.join(parts))
            print(f'  R{ri}: ' + '  ||  '.join(cells))
    elif tag == 'sectPr':
        print('SECTPR')
