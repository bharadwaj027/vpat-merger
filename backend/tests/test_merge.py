import copy
import io

import pytest
from conftest import FIXTURES, cell_platform_texts
from docx import Document
from docx.oxml.ns import qn

from vpat_merger import Platform, merge_vpats
from vpat_merger.docx_utils import el_text, norm_space
from vpat_merger.matching import union_order
from vpat_merger.models import MISSING_MARKER
from vpat_merger.parser import parse_vpat

D, I, A = Platform.DESKTOP, Platform.IOS, Platform.ANDROID


def run(files):
    result = merge_vpats(files)
    assert result.docx is not None, result.report
    return result


def issues(report, category=None, criterion=None):
    return [i for i in report["issues"]
            if (category is None or i["category"] == category) and (criterion is None or i["criterion"] == criterion)]


def section(texts, label):
    """Slice of merged-cell paragraphs belonging to one platform label."""
    labels = ["Desktop Web:", "iOS Mobile Web:", "Android Mobile Web:"]
    start = next(i for i, t in enumerate(texts) if t.startswith(label))
    end = next((i for i, t in enumerate(texts[start + 1:], start + 1) if any(t.startswith(l) for l in labels)),
               len(texts))
    return [t for t in texts[start:end] if t]


# 1. Criteria present in all three sources --------------------------------------------------
def test_all_three_sources_merge_cleanly(sample_files):
    result = run(sample_files)
    rep = result.report
    assert rep["combined"]["uniqueCriteria"] == 55
    assert rep["combined"]["inAllSources"] == 55
    assert rep["outputValid"], [c for c in rep["checks"] if not c["passed"]]
    assert rep["issueCounts"]["error"] == 0
    assert result.filename == "TestCo-IdentityDashboard-Combined-VPAT.docx"


# 2. Criteria present in only two sources ---------------------------------------------------
def test_criterion_missing_from_one_source(sample_files, mutate):
    sample_files[A] = mutate(A).remove_row("2.4.11").bytes()
    rep = run(sample_files).report
    assert rep["combined"]["uniqueCriteria"] == 55
    missing = {m["key"]: m["missing"] for m in rep["combined"]["missingFromSome"]}
    assert missing == {"2.4.11": ["Android Mobile Web"]}
    assert issues(rep, "missing-criterion", "2.4.11")
    assert rep["outputValid"]


def test_missing_criterion_is_marked_not_invented(sample_files, mutate):
    sample_files[A] = mutate(A).remove_row("2.4.11").bytes()
    result = run(sample_files)
    conf = cell_platform_texts(result.docx, "2.4.11", 1)
    rem = cell_platform_texts(result.docx, "2.4.11", 2)
    assert section(conf, "Android Mobile Web:") == [f"Android Mobile Web: {MISSING_MARKER}"]
    assert section(rem, "Android Mobile Web:") == [f"Android Mobile Web: {MISSING_MARKER}"]
    # Never silently treated as Not Applicable.
    assert "Not Applicable" not in " ".join(section(conf, "Android Mobile Web:"))


# 3. Criteria present in only one source ----------------------------------------------------
def test_criterion_in_only_one_source(sample_files, mutate):
    sample_files[I] = mutate(I).remove_row("1.4.13").bytes()
    sample_files[A] = mutate(A).remove_row("1.4.13").bytes()
    result = run(sample_files)
    rep = result.report
    missing = {m["key"]: m["missing"] for m in rep["combined"]["missingFromSome"]}
    assert missing == {"1.4.13": ["iOS Mobile Web", "Android Mobile Web"]}
    conf = cell_platform_texts(result.docx, "1.4.13", 1)
    assert section(conf, "iOS Mobile Web:") == [f"iOS Mobile Web: {MISSING_MARKER}"]
    assert section(conf, "Desktop Web:")[0] != f"Desktop Web: {MISSING_MARKER}"
    assert rep["outputValid"]


def test_criterion_only_in_mobile_sources_is_still_included(sample_files, mutate):
    sample_files[D] = mutate(D).remove_row("1.3.4").bytes()
    result = run(sample_files)
    keys = [m["key"] for m in result.report["matrix"]]
    assert keys.count("1.3.4") == 1
    # Logical order is kept even though the base (Desktop) document lacks the row.
    assert keys.index("1.3.4") == keys.index("1.2.5") + 1
    assert result.report["outputValid"]


# 4. Different conformance results for the same criterion ------------------------------------
def test_different_conformance_values_kept_per_platform(sample_files):
    result = run(sample_files)
    conf = cell_platform_texts(result.docx, "1.1.1", 1)
    assert section(conf, "Desktop Web:") == ["Desktop Web: Does Not Support"]
    assert section(conf, "iOS Mobile Web:") == ["iOS Mobile Web: Partially Supports"]
    assert section(conf, "Android Mobile Web:") == ["Android Mobile Web: Partially Supports"]


def test_no_combined_or_overall_level_is_computed(sample_files, mutate):
    sample_files[I] = mutate(I).set_cell_text("2.1.1", 1, "Supports").bytes()
    sample_files[A] = mutate(A).set_cell_text("2.1.1", 1, "Not Applicable").bytes()
    result = run(sample_files)
    row = next(r for r in result.report["matrix"] if r["key"] == "2.1.1")
    desktop = parse_vpat(*sample_files[D], D)
    src = next(r for t in desktop.tables for r in t.rows if r.key == "2.1.1")
    assert row["results"] == {"desktop": src.conformance, "ios": "Supports", "android": "Not Applicable"}
    conf = cell_platform_texts(result.docx, "2.1.1", 1)
    assert section(conf, "Android Mobile Web:") == ["Android Mobile Web: Not Applicable"]


# 5. Different remarks for the same criterion ------------------------------------------------
def test_remarks_stay_under_their_own_platform(sample_files, mutate):
    marker = "IOS-ONLY remark about the date picker"
    sample_files[I] = mutate(I).set_cell_text("1.3.1", 2, marker).bytes()
    result = run(sample_files)
    rem = cell_platform_texts(result.docx, "1.3.1", 2)
    assert section(rem, "iOS Mobile Web:") == [f"iOS Mobile Web: {marker}"]
    assert marker not in " ".join(section(rem, "Desktop Web:") + section(rem, "Android Mobile Web:"))
    assert result.report["outputValid"]


# 6. Empty remarks ----------------------------------------------------------------------------
def test_empty_remarks_flagged_and_not_backfilled(sample_files, mutate):
    sample_files[A] = mutate(A).set_cell_text("2.1.2", 2, "").bytes()
    result = run(sample_files)
    rem = cell_platform_texts(result.docx, "2.1.2", 2)
    assert section(rem, "Android Mobile Web:") == ["Android Mobile Web:"]
    assert issues(result.report, "empty-remarks", "2.1.2")
    assert result.report["outputValid"]


# 7. Missing / invalid conformance levels ----------------------------------------------------
def test_missing_conformance_flagged(sample_files, mutate):
    sample_files[I] = mutate(I).set_cell_text("1.4.1", 1, "").bytes()
    result = run(sample_files)
    found = issues(result.report, "missing-conformance", "1.4.1")
    assert found and found[0]["severity"] == "error" and found[0]["platform"] == "ios"
    conf = cell_platform_texts(result.docx, "1.4.1", 1)
    assert section(conf, "iOS Mobile Web:") == [f"iOS Mobile Web: {MISSING_MARKER}"]


def test_invalid_conformance_kept_verbatim_and_flagged(sample_files, mutate):
    sample_files[A] = mutate(A).set_cell_text("2.4.2", 1, "Mostly Supports").bytes()
    result = run(sample_files)
    assert issues(result.report, "invalid-conformance", "2.4.2")
    conf = cell_platform_texts(result.docx, "2.4.2", 1)
    assert section(conf, "Android Mobile Web:") == ["Android Mobile Web: Mostly Supports"]


# 8. Multi-paragraph remarks -----------------------------------------------------------------
def test_multi_paragraph_remarks_keep_order_and_breaks(sample_files):
    result = run(sample_files)
    rem = cell_platform_texts(result.docx, "1.1.1", 2)
    for p in (D, I, A):
        src = parse_vpat(*sample_files[p], p)
        row = next(r for t in src.tables for r in t.rows if r.key == "1.1.1")
        got = section(rem, f"{p.label}:")
        got[0] = got[0][len(p.label) + 2:]
        assert got == row.remarks_text
        assert len(got) > 1


# 9. Bullet lists and long remarks -----------------------------------------------------------
def test_bullets_and_longest_remarks_preserved(sample_files):
    result = run(sample_files)
    out = Document(io.BytesIO(result.docx))
    for p in (D, I, A):
        src = parse_vpat(*sample_files[p], p)
        longest = max((r for t in src.tables for r in t.rows), key=lambda r: sum(map(len, r.remarks_text)))
        assert sum(map(len, longest.remarks_text)) > 500
        rem = cell_platform_texts(result.docx, longest.key, 2)
        got = section(rem, f"{p.label}:")
        got[0] = got[0][len(p.label) + 2:]
        assert got == longest.remarks_text
    # Bulleted paragraphs keep their numbering properties in the output.
    tbl = next(t for t in out.tables if "Criteria" in t.rows[0].cells[0].text)
    cell = tbl.rows[1].cells[2]._tc
    assert sum(1 for p in cell.iter(qn("w:p")) if p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None) >= 9


# 10. Duplicate or malformed criteria --------------------------------------------------------
def test_duplicate_criterion_is_flagged_and_merged_once(sample_files, mutate):
    sample_files[D] = mutate(D).duplicate_row("1.4.1", "Supports").bytes()
    result = run(sample_files)
    dup = issues(result.report, "duplicate-criterion", "1.4.1")
    assert dup and dup[0]["severity"] == "error"
    keys = [m["key"] for m in result.report["matrix"]]
    assert keys.count("1.4.1") == 1


def test_malformed_criterion_is_kept_and_flagged(sample_files, mutate):
    sample_files[I] = mutate(I).set_cell_text("3.1.1", 0, "Language check without identifier").bytes()
    result = run(sample_files)
    rep = result.report
    assert issues(rep, "malformed-criterion")
    titles = [m["title"] for m in rep["matrix"]]
    assert "Language check without identifier" in titles
    missing = {m["key"]: m["missing"] for m in rep["combined"]["missingFromSome"]}
    assert missing["3.1.1"] == ["iOS Mobile Web"]
    assert rep["outputValid"]


def test_non_vpat_document_is_rejected(tmp_path):
    from docx import Document as NewDoc

    d = NewDoc()
    d.add_paragraph("Just a memo")
    buf = io.BytesIO()
    d.save(buf)
    result = merge_vpats({D: (buf.getvalue(), "memo.docx"),
                          I: (buf.getvalue(), "memo.docx"),
                          A: (buf.getvalue(), "memo.docx")})
    assert result.docx is None
    assert "conformance table" in result.report["errors"][0]


def test_non_docx_is_rejected(sample_files):
    sample_files[I] = (b"%PDF-1.7 not a word file", "ios.pdf")
    result = merge_vpats(sample_files)
    assert result.docx is None
    assert any("not a .docx" in e for e in result.report["errors"])


# 11. Source metadata conflicts ---------------------------------------------------------------
def test_metadata_conflicts_are_labeled_and_flagged(sample_files, mutate):
    sample_files[I] = mutate(I).set_paragraph_after_heading("Report Date:", "September 30, 2026").bytes()
    result = run(sample_files)
    doc = Document(io.BytesIO(result.docx))
    texts = [p.text for p in doc.paragraphs]
    i = texts.index("Report Date:")
    assert texts[i + 1:i + 4] == ["Desktop Web: October 28, 2026", "iOS Mobile Web: September 30, 2026",
                                  "Android Mobile Web: October 28, 2026"]
    conflicts = [x["message"] for x in issues(result.report, "metadata-conflict")]
    assert any("Report Date" in m for m in conflicts)
    assert any("Name of Product" in m for m in conflicts)


def test_evaluation_methods_keep_assistive_tech_per_platform(sample_files):
    doc = Document(io.BytesIO(run(sample_files).docx))
    texts = [p.text for p in doc.paragraphs]
    start = texts.index("Evaluation Methods Used:")
    block = texts[start + 1:texts.index("Scope of Evaluation")]
    desktop = next(t for t in block if t.startswith("Desktop Web:"))
    ios = next(t for t in block if t.startswith("iOS Mobile Web:"))
    android = next(t for t in block if t.startswith("Android Mobile Web:"))
    assert "NVDA" in desktop and "NVDA" not in ios + android
    assert "Voiceover" in ios and "Voiceover" not in desktop + android
    assert "Talkback" in android and "Talkback" not in desktop + ios


def test_scope_keeps_every_page_once_in_two_columns(sample_files):
    result = run(sample_files)
    doc = Document(io.BytesIO(result.docx))
    scope = doc.tables[0]
    assert [c.text for c in scope.rows[0].cells] == ["Page Title", "URL"]
    rows = [(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in scope.rows[1:]]
    assert all(len(r.cells) == 2 for r in scope.rows)
    assert len(rows) == len(set(rows))
    expected = set()
    for p in (D, I, A):
        src = Document(io.BytesIO(sample_files[p][0])).tables[0]
        expected |= {(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in src.rows[1:]}
    assert set(rows) == expected
    assert ("Email Link", "https://verify.testco.example/handoff") in rows  # Desktop-only page is kept


# 12. Output document integrity --------------------------------------------------------------
def test_output_document_integrity(sample_files):
    result = run(sample_files)
    doc = Document(io.BytesIO(result.docx))
    assert doc.paragraphs[1].text == "TestCo-IdentityDashboard Desktop and Mobile Web Accessibility Conformance Report"
    crit_tables = [t for t in doc.tables if "Criteria" in t.rows[0].cells[0].text]
    assert [len(t.rows) - 1 for t in crit_tables] == [31, 24]
    for t in crit_tables:
        assert [c.text for c in t.rows[0].cells] == ["Criteria", "Conformance Level", "Remarks and Explanations"]
        assert t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None  # header repeats across pages
        for r in t.rows[1:]:
            assert r._tr.trPr is None or r._tr.trPr.find(qn("w:cantSplit")) is None  # long rows may break
            for c in r.cells:
                assert norm_space(c.text)
            label_run = r.cells[1].paragraphs[0].runs[0]
            assert label_run.text == "Desktop Web:" and label_run.bold
    # Hyperlinks (criteria links + scope URLs) resolve after import from the iOS/Android packages.
    for node in doc.element.body.iter(qn("w:hyperlink")):
        rid = node.get(qn("r:id"))
        if rid:
            assert rid in doc.part.rels
    # Only the uploaded sources' content may appear: nothing from other documents (e.g. the native fixtures).
    body = norm_space(el_text(doc.element.body))
    assert "SampleApp" not in body
    assert "Not Applicable" in body and result.report["outputValid"]


def test_not_applicable_is_never_changed(sample_files):
    result = run(sample_files)
    for row in result.report["matrix"]:
        for p in (D, I, A):
            src = parse_vpat(*sample_files[p], p)
            want = next(r.conformance for t in src.tables for r in t.rows if r.key == row["key"])
            assert row["results"][p.value] == want


@pytest.mark.parametrize("seqs,expected", [
    ([["a", "b", "d"], ["a", "c", "d"]], ["a", "b", "c", "d"]),
    ([["1.2.4", "1.2.5", "1.4.3"], ["1.2.4", "1.3.4", "1.4.3"]], ["1.2.4", "1.2.5", "1.3.4", "1.4.3"]),
    ([["1.4.3"], ["1.3.4", "1.4.3"]], ["1.3.4", "1.4.3"]),
])
def test_union_order(seqs, expected):
    assert union_order(seqs) == expected


# Two-platform merges -------------------------------------------------------------------------
@pytest.mark.parametrize("pair,title_scope", [
    ((D, I), "Desktop and Mobile Web"),
    ((D, A), "Desktop and Mobile Web"),
    ((I, A), "Mobile Web"),
])
def test_two_platform_merge(sample_files, pair, title_scope):
    files = {p: sample_files[p] for p in pair}
    result = run(files)
    rep = result.report
    assert rep["outputValid"], [c for c in rep["checks"] if not c["passed"]]
    assert rep["platforms"] == [p.value for p in pair]
    assert rep["combined"]["uniqueCriteria"] == 55
    assert rep["combined"]["inAllSources"] == 55
    assert all(set(row["results"]) == {p.value for p in pair} for row in rep["matrix"])

    doc = Document(io.BytesIO(result.docx))
    assert doc.paragraphs[1].text == f"TestCo-IdentityDashboard {title_scope} Accessibility Conformance Report"
    absent = next(p for p in (D, I, A) if p not in pair)
    body = norm_space(el_text(doc.element.body))
    # The platform that was not uploaded is never labeled or reported as missing.
    assert f"{absent.label}:" not in body
    assert MISSING_MARKER not in body
    conf = cell_platform_texts(result.docx, "1.1.1", 1)
    assert [t for t in conf if t] == [section(conf, f"{p.label}:")[0] for p in pair]


def test_two_platform_missing_criterion(sample_files, mutate):
    files = {I: mutate(I).remove_row("2.5.8").bytes(), A: sample_files[A]}
    rep = run(files).report
    assert {m["key"]: m["missing"] for m in rep["combined"]["missingFromSome"]} == {"2.5.8": ["iOS Mobile Web"]}
    assert rep["outputValid"]


def test_single_source_is_rejected(sample_files):
    result = merge_vpats({I: sample_files[I]})
    assert result.docx is None
    assert "at least 2" in result.report["errors"][0]


# Native app VPATs ----------------------------------------------------------------------------
NATIVE = {
    Platform.IOS_NATIVE: FIXTURES / "native" / "ios-native.docx",
    Platform.ANDROID_NATIVE: FIXTURES / "native" / "android-native.docx",
}


@pytest.fixture
def native_files():
    return {p: (path.read_bytes(), path.name) for p, path in NATIVE.items()}


def test_native_labels_in_both_columns(native_files):
    result = run(native_files)
    rep = result.report
    assert rep["outputValid"], [c for c in rep["checks"] if not c["passed"]]
    assert rep["platformLabels"] == {"ios_native": "Native iOS", "android_native": "Native Android"}
    for col in (1, 2):
        texts = [t for t in cell_platform_texts(result.docx, "1.1.1", col) if t]
        assert texts[0].startswith("Native iOS: ")
        assert any(t.startswith("Native Android: ") for t in texts)
        assert not any("Mobile Web:" in t for t in texts)
    doc = Document(io.BytesIO(result.docx))
    assert doc.paragraphs[1].text == "SampleApp Native iOS and Android Accessibility Conformance Report"


def test_native_single_column_scope_is_merged(native_files):
    doc = Document(io.BytesIO(run(native_files).docx))
    scope = doc.tables[0]
    assert [c.text for c in scope.rows[0].cells] == ["Page Title"]
    rows = [r.cells[0].text.strip() for r in scope.rows[1:]]
    assert len(rows) == len(set(rows))
    for p in native_files:
        src = Document(io.BytesIO(native_files[p][0])).tables[0]
        assert {r.cells[0].text.strip() for r in src.rows[1:]} <= set(rows)


def test_native_platforms_are_detected(native_files, sample_files):
    from vpat_merger.parser import guess_platform

    expected = {**{p: p for p in native_files}, **{p: p for p in sample_files}}
    for p, (data, name) in {**native_files, **sample_files}.items():
        assert guess_platform(parse_vpat(data, name, Platform.IOS)) is expected[p]


def test_web_and_native_mixed(sample_files, native_files):
    files = {D: sample_files[D], Platform.IOS_NATIVE: native_files[Platform.IOS_NATIVE]}
    result = run(files)
    assert result.report["platforms"] == ["desktop", "ios_native"]
    conf = cell_platform_texts(result.docx, "1.1.1", 1)
    assert [t.split(":")[0] for t in conf if t] == ["Desktop Web", "Native iOS"]
    assert result.report["outputValid"]


# Bold platform prefixes ----------------------------------------------------------------------
def _label_runs(docx_bytes, labels):
    doc = Document(io.BytesIO(docx_bytes))
    for p in doc.element.body.iter(qn("w:p")):
        if norm_space(el_text(p)).startswith(tuple(f"{l}:" for l in labels)):
            yield next(r for r in p.iter(qn("w:r")) if norm_space(el_text(r)))


@pytest.mark.parametrize("which", ["web", "native"])
def test_every_prefix_is_bold_with_schema_ordered_formatting(sample_files, native_files, which):
    files = sample_files if which == "web" else native_files
    result = run(files)
    labels = list(result.report["platformLabels"].values())
    runs = list(_label_runs(result.docx, labels))
    assert len(runs) > 100
    for r in runs:
        rpr = r.find(qn("w:rPr"))
        assert rpr is not None and rpr.find(qn("w:b")) is not None
        names = [c.tag.split("}")[1] for c in rpr]
        # Bold must come after the font and before colour/size, as the WordprocessingML schema requires.
        assert names.index("b") < names.index("sz")
        if "rFonts" in names:
            assert names.index("rFonts") < names.index("b")
        assert "shd" not in names and "rStyle" not in names
        assert norm_space(el_text(r)) in {f"{l}:" for l in labels}
    checks = {c["name"]: c["passed"] for c in result.report["checks"]}
    assert checks["Every platform prefix is bold"]
    assert checks["Run formatting is valid for all Word viewers (bold labels display everywhere)"]


def test_source_rpr_order_is_repaired_without_losing_formatting(sample_files):
    from vpat_merger.docx_utils import _rpr_rank

    src = Document(io.BytesIO(sample_files[D][0]))
    assert any([_rpr_rank(c) for c in rpr] != sorted(_rpr_rank(c) for c in rpr)
               for rpr in src.element.body.iter(qn("w:rPr")))  # the fixtures really are out of order
    doc = Document(io.BytesIO(run(sample_files).docx))
    # Source bold/underline/colour runs keep their formatting after reordering.
    def styled(d):
        return sum(1 for r in d.element.body.iter(qn("w:r"))
                   if r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:u")) is not None)
    assert styled(doc) >= styled(src)
    for rpr in doc.element.body.iter(qn("w:rPr")):
        ranks = [_rpr_rank(c) for c in rpr]
        assert ranks == sorted(ranks)


# Scope: pages present on one platform but not another ----------------------------------------
def test_desktop_only_pages_are_flagged(sample_files):
    rep = run(sample_files).report
    diffs = {d["page"]: d for d in rep["scopeDifferences"]}
    assert set(diffs) == {"Continue to your mobile device", "Email Link", "Others options"}
    for d in diffs.values():
        assert d["presentIn"] == ["Desktop Web"]
        assert d["missingFrom"] == ["iOS Mobile Web", "Android Mobile Web"]
    flagged = [i for i in rep["issues"] if i["category"] == "scope-page-missing"]
    assert len(flagged) == 3 and all(i["severity"] == "warning" for i in flagged)
    assert any("'Email Link' is evaluated in Desktop Web but is not listed in the iOS Mobile Web and Android Mobile Web"
               in i["message"] for i in flagged)


def test_same_page_with_different_url_or_spelling_is_not_flagged(sample_files):
    pages = {d["page"] for d in run(sample_files).report["scopeDifferences"]}
    assert "Allow Camera access" not in pages  # different URL on Android
    assert not {"Project wide", "Project-wide"} & pages  # hyphen vs space


def test_mobile_only_page_is_flagged(sample_files):
    doc = Document(io.BytesIO(sample_files[I][0]))
    scope = doc.tables[0]
    new = copy.deepcopy(scope.rows[1]._tr)
    scope._tbl.append(new)
    row = scope.rows[-1]
    for cell, text in zip(row.cells, ["Mobile menu drawer", "https://verify.testco.example/menu"]):
        for t in cell._tc.iter(qn("w:t")):
            t.text = ""
        next(cell._tc.iter(qn("w:t"))).text = text
    buf = io.BytesIO()
    doc.save(buf)
    sample_files[I] = (buf.getvalue(), sample_files[I][1])
    rep = run(sample_files).report
    d = next(d for d in rep["scopeDifferences"] if d["page"] == "Mobile menu drawer")
    assert d["presentIn"] == ["iOS Mobile Web"]
    assert d["missingFrom"] == ["Desktop Web", "Android Mobile Web"]


def test_identical_scope_has_no_flags(sample_files):
    rep = run({D: sample_files[D], I: sample_files[D]}).report
    assert rep["scopeDifferences"] == []
    assert not [i for i in rep["issues"] if i["category"] == "scope-page-missing"]


def _scope_rows(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    return doc, doc.tables[0].rows[1:]


def test_flagged_pages_are_highlighted_and_commented_in_document(sample_files):
    result = run(sample_files)
    doc, rows = _scope_rows(result.docx)
    flagged = {"Continue to your mobile device", "Email Link", "Others options"}
    comments = {c.comment_id: c for c in doc.comments}
    assert len(comments) == 3
    for row in rows:
        title = row.cells[0].text.strip()
        tr = row._tr
        shaded = all(tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}") is not None for tc in tr.findall(qn("w:tc")))
        highlighted = [r.find(f"{qn('w:rPr')}/{qn('w:highlight')}") is not None
                       for r in tr.iter(qn("w:r")) if norm_space(el_text(r))]
        ids = [int(c.get(qn("w:id"))) for c in tr.iter(qn("w:commentRangeStart"))]
        if title in flagged:
            assert shaded and all(highlighted)
            assert len(ids) == 1
            c = comments[ids[0]]
            assert c.author == "VPAT Merger"
            assert c.text == ("This page is not present in the iOS Mobile Web and Android Mobile Web VPATs. "
                              "It was evaluated on Desktop Web only.")
        else:
            assert not shaded and not any(highlighted) and not ids
    checks = {c["name"]: c["passed"] for c in result.report["checks"]}
    assert checks["Pages missing from a platform are highlighted and commented in the scope table"]


def test_mobile_only_page_comment_names_desktop(sample_files):
    doc = Document(io.BytesIO(sample_files[I][0]))
    scope = doc.tables[0]
    scope._tbl.append(copy.deepcopy(scope.rows[1]._tr))
    for cell, text in zip(scope.rows[-1].cells, ["Mobile menu drawer", "https://verify.testco.example/menu"]):
        for t in cell._tc.iter(qn("w:t")):
            t.text = ""
        next(cell._tc.iter(qn("w:t"))).text = text
    buf = io.BytesIO()
    doc.save(buf)
    files = {D: sample_files[D], I: (buf.getvalue(), sample_files[I][1])}
    out, _ = _scope_rows(run(files).docx)
    texts = [c.text for c in out.comments]
    assert "This page is not present in the Desktop Web VPAT. It was evaluated on iOS Mobile Web only." in texts


def test_no_comments_when_scope_matches(sample_files):
    doc, rows = _scope_rows(run({D: sample_files[D], I: sample_files[D]}).docx)
    assert len(list(doc.comments)) == 0
    assert not any(tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}") is not None
                   for row in rows for tc in row._tr.findall(qn("w:tc")))
