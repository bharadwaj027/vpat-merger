import copy
import io
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from vpat_merger import Platform  # noqa: E402
from vpat_merger.docx_utils import el_text, norm_space  # noqa: E402

FIXTURES = ROOT / "fixtures"
FILES = {
    Platform.DESKTOP: FIXTURES / "web.docx",
    Platform.IOS: FIXTURES / "ios.docx",
    Platform.ANDROID: FIXTURES / "android.docx",
}


def pytest_collection_modifyitems(config, items):
    """The sample VPATs are real client reports and may be absent; skip tests that need them."""
    if all(path.exists() for path in FILES.values()):
        return
    skip = pytest.mark.skip(reason="fixture VPATs not present in fixtures/ (client documents are not distributed)")
    for item in items:
        if "union_order" not in item.name:
            item.add_marker(skip)


@pytest.fixture
def sample_files():
    return {p: (path.read_bytes(), path.name) for p, path in FILES.items()}


class Mutator:
    """Edits a copy of a fixture VPAT to build edge-case inputs."""

    def __init__(self, platform: Platform):
        self.platform = platform
        self.doc = Document(str(FILES[platform]))

    def criteria_tables(self):
        return [t for t in self.doc.tables if "Criteria" in t.rows[0].cells[0].text]

    def row(self, sc_id: str):
        for t in self.criteria_tables():
            for r in t.rows[1:]:
                if r.cells[0].text.strip().startswith(sc_id + " "):
                    return r
        raise KeyError(sc_id)

    def remove_row(self, sc_id: str):
        tr = self.row(sc_id)._tr
        tr.getparent().remove(tr)
        return self

    def set_cell_text(self, sc_id: str, col: int, text: str):
        tc = self.row(sc_id).cells[col]._tc
        paras = tc.findall(qn("w:p"))
        for p in paras[1:]:
            tc.remove(p)
        for tbl in tc.findall(qn("w:tbl")):
            tc.remove(tbl)
        p = paras[0]
        first = next(iter(p.iter(qn("w:r"))), None)
        rpr = first.find(qn("w:rPr")) if first is not None else None
        for child in list(p):
            if child.tag != qn("w:pPr"):
                p.remove(child)
        if text:
            r = _run(text)
            if rpr is not None:
                r.insert(0, copy.deepcopy(rpr))
            p.append(r)
        return self

    def duplicate_row(self, sc_id: str, conformance: str):
        tr = self.row(sc_id)._tr
        dup = copy.deepcopy(tr)
        tr.addnext(dup)
        cell = dup.findall(qn("w:tc"))[1]
        for t in cell.iter(qn("w:t")):
            t.text = ""
        cell.find(qn("w:p")).append(_run(conformance))
        return self

    def set_paragraph_after_heading(self, heading: str, text: str):
        paras = self.doc.paragraphs
        for i, p in enumerate(paras):
            if p.text.strip() == heading:
                target = next(q for q in paras[i + 1:] if q.text.strip())
                for r in target.runs[1:]:
                    r._r.getparent().remove(r._r)
                target.runs[0].text = text
                return self
        raise KeyError(heading)

    def bytes(self):
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue(), FILES[self.platform].name


def _run(text):
    from docx.oxml import OxmlElement

    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    return r


@pytest.fixture
def mutate():
    return Mutator


def cell_platform_texts(docx_bytes: bytes, sc_id: str, col: int) -> list[str]:
    """Paragraph texts of one merged cell, for assertions."""
    doc = Document(io.BytesIO(docx_bytes))
    for t in doc.tables:
        if "Criteria" not in t.rows[0].cells[0].text:
            continue
        for r in t.rows[1:]:
            if norm_space(r.cells[0].text).startswith(sc_id + " "):
                return [norm_space(el_text(p._p)) for p in r.cells[col].paragraphs]
    raise KeyError(sc_id)
