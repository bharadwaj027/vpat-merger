"""Build the combined Word document.

The Desktop Web source is used as the base package, so its styles, numbering,
headers/footers, section breaks and logo carry over. Content from the iOS and
Android sources is deep-copied in through an Importer that re-creates their
hyperlink/image relationships and numbering ids.
"""
from __future__ import annotations

import copy
import io
import re
from typing import Optional, Sequence

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

from .docx_utils import (
    Importer,
    cell_blocks,
    el_text,
    empty_paragraph_like,
    first_run_rpr,
    insert_runs_at_start,
    is_list_para,
    is_pure_sectpr_para,
    make_run,
    norm_space,
    normalize_all_rpr,
    para_style_id,
    paragraph_from_template,
)
from .matching import MergedCriterion, MergedTable, union_order
from .models import MISSING_MARKER, Issue, Platform, Section, SourceVpat, active_platforms
from .scope import page_key, scope_comment

SCOPE_HIGHLIGHT = "yellow"
SCOPE_SHADING = "FFFF00"
COMMENT_AUTHOR = "VPAT Merger"

# CT_TcPr children that must precede <w:shd> in schema order.
_TCPR_BEFORE_SHD = {"cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders"}


def shade_cell(tc, fill: str):
    tcpr = tc.find(qn("w:tcPr"))
    if tcpr is None:
        tcpr = OxmlElement("w:tcPr")
        tc.insert(0, tcpr)
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pos = sum(1 for c in tcpr if c.tag.split("}")[1] in _TCPR_BEFORE_SHD)
    tcpr.insert(pos, shd)

def combined_product_name(sources: dict[Platform, SourceVpat]) -> str:
    names = [sources[p].product_name for p in active_platforms(sources) if sources[p].product_name]
    if not names:
        return "Combined"
    if len(set(names)) == 1:
        return names[0]
    prefix = names[0]
    for n in names[1:]:
        i = 0
        while i < min(len(prefix), len(n)) and prefix[i] == n[i]:
            i += 1
        prefix = prefix[:i]
    # Cut back to a token boundary so "Product-M" (from "Product-MWEB...") does not leak a partial token.
    cuts_token = any(len(n) > len(prefix) and n[len(prefix)].isalnum() for n in names)
    if prefix and prefix[-1].isalnum() and cuts_token:
        prefix = re.sub(r"[^\-_ .,]*$", "", prefix)
    prefix = prefix.strip(" -_.,")
    # Drop trailing platform words ("SampleApp-Native Mobile" → "SampleApp"); the title names the platforms itself.
    platform_word = re.compile(r"[\s\-_.,]*\b(native|mobile|web|mweb|apps?|ios|android|desktop|audit)$", re.IGNORECASE)
    while platform_word.search(prefix) and len(platform_word.sub("", prefix)) >= 3:
        prefix = platform_word.sub("", prefix)
    return prefix if len(prefix) >= 3 else names[0]


def report_scope(platforms: list[Platform]) -> str:
    """Title wording for the combined report, e.g. 'Desktop and Mobile Web' or 'Native iOS and Android'."""
    parts = []
    mobile_web = [p for p in platforms if p in (Platform.IOS, Platform.ANDROID)]
    if Platform.DESKTOP in platforms:
        parts.append("Desktop and Mobile Web" if mobile_web else "Desktop Web")
    elif mobile_web:
        parts.append("Mobile Web" if len(mobile_web) > 1 else mobile_web[0].label)
    native = [p for p in platforms if p.is_native]
    if len(native) > 1:
        parts.append("Native iOS and Android")
    elif native:
        parts.append(native[0].label)
    return ", ".join(parts)


def _signature(blocks) -> list[tuple]:
    sig = []
    for b in blocks:
        if is_pure_sectpr_para(b):
            continue
        text = norm_space(el_text(b))
        if b.tag == qn("w:p") and not text:
            continue
        sig.append((b.tag, para_style_id(b) if b.tag == qn("w:p") else "", is_list_para(b), text))
    return sig


def _content(blocks) -> list:
    """Blocks without section-break carriers and leading/trailing blank paragraphs."""
    out = [b for b in blocks if not is_pure_sectpr_para(b)]
    while out and out[0].tag == qn("w:p") and not norm_space(el_text(out[0])) and out[0].find(f".//{qn('w:drawing')}") is None:
        out.pop(0)
    while out and out[-1].tag == qn("w:p") and not norm_space(el_text(out[-1])) and out[-1].find(f".//{qn('w:drawing')}") is None:
        out.pop()
    return out


def set_paragraph_text(p, text: str):
    rpr = first_run_rpr(p)
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    p.append(make_run(text, rpr))


class Generator:
    def __init__(self, sources: dict[Platform, SourceVpat], merged_tables: list[MergedTable], issues: list[Issue],
                 scope_differences: Sequence[dict] = ()):
        self.sources = sources
        self.merged_tables = {mt.key: mt for mt in merged_tables}
        self.issues = issues
        self.scope_differences = list(scope_differences)
        self._scope_tables: list = []
        self.flagged_rows = 0
        self.platforms = active_platforms(sources)
        # The first uploaded platform (Desktop when present) supplies styles, headers/footers and page setup.
        self.base_platform = self.platforms[0]
        self.base = sources[self.base_platform]
        self.doc = self.base.document
        self.imp = {p: Importer(sources[p].document, self.doc) for p in self.platforms}
        self.product = combined_product_name(sources)
        self.title = f"{self.product} {report_scope(self.platforms)} Accessibility Conformance Report"
        self._list_template = self._find_list_template()

    # ---------- helpers ----------
    def _find_list_template(self):
        for p in self.doc.element.body.iter(qn("w:p")):
            if is_list_para(p) and norm_space(el_text(p)):
                return p
        return None

    def label_runs(self, platform: Platform, rpr):
        return [make_run(f"{platform.label}:", rpr, bold=True), make_run(" ", rpr, bold=False)]

    def label_paragraph(self, platform: Platform, template_p, trailing: Optional[str] = None, italic=False):
        p = empty_paragraph_like(template_p) if template_p is not None else OxmlElement("w:p")
        rpr = first_run_rpr(template_p) if template_p is not None else None
        p.append(make_run(f"{platform.label}:", rpr, bold=True))
        if trailing:
            p.append(make_run(" ", rpr, bold=False))
            p.append(make_run(trailing, rpr, bold=False, italic=italic))
        return p

    def labeled_blocks(self, platform: Platform, blocks: list, template_p=None) -> list:
        """Platform label followed by that platform's (already imported) blocks."""
        if not blocks:
            return [self.label_paragraph(platform, template_p)]
        first = blocks[0]
        if first.tag == qn("w:p") and not is_list_para(first) and norm_space(el_text(first)):
            insert_runs_at_start(first, self.label_runs(platform, first_run_rpr(first)))
            return blocks
        tmpl = template_p if template_p is not None else next(
            (b for b in blocks if b.tag == qn("w:p") and not is_list_para(b)), None)
        return [self.label_paragraph(platform, tmpl)] + blocks

    def merge_simple(self, key: str, parts: dict[Platform, Optional[list]], what: str) -> list:
        """Keep shared content once; show differing content per platform with labels."""
        present = [p for p in self.platforms if parts.get(p) is not None]
        sigs = {p: _signature(parts[p]) for p in present}
        base_blocks = parts.get(self.base_platform)
        if len(present) == len(self.platforms) and all(sigs[p] == sigs[self.base_platform] for p in present):
            return list(base_blocks)
        if all(not sigs[p] for p in present) and base_blocks is not None:
            return list(base_blocks)
        missing = [p for p in self.platforms if p not in present]
        self.issues.append(Issue(
            "info" if not missing else "warning", "metadata-conflict",
            f"'{what}' differs between the source VPATs"
            + (f" (missing from {', '.join(p.label for p in missing)})" if missing else "")
            + ". Each platform's content is kept under its own label; review the wording.",
        ))
        template_p = None
        for p in present:
            template_p = next((b for b in _content(parts[p]) if b.tag == qn("w:p") and not is_list_para(b)), None)
            if template_p is not None:
                break
        out = []
        for p in self.platforms:
            if parts.get(p) is None:
                out.append(self.label_paragraph(p, template_p, MISSING_MARKER, italic=True))
                continue
            imported = [self.imp[p](b) for b in _content(parts[p])]
            out.extend(self.labeled_blocks(p, imported, template_p))
        if base_blocks is not None:
            out.extend(b for b in base_blocks if is_pure_sectpr_para(b))
        return out

    # ---------- section handlers ----------
    def build(self) -> bytes:
        section_keys = union_order([[s.key for s in self.sources[p].sections] for p in self.platforms])
        new_body: list = []
        for key in section_keys:
            secs = {p: self.sources[p].section(key) for p in self.platforms}
            present = [p for p in self.platforms if secs[p] is not None]
            first = secs[present[0]]
            if first.heading_el is not None:
                if secs[self.base_platform] is not None:
                    heading = secs[self.base_platform].heading_el
                else:
                    heading = self.imp[present[0]](first.heading_el)
                self._retitle_heading(heading, key, first)
                new_body.append(heading)
            new_body.extend(self.handle_section(key, secs))

        body = self.doc.element.body
        final_sectpr = body.find(qn("w:sectPr"))
        for child in list(body):
            if child is not final_sectpr:
                body.remove(child)
        for el in new_body:
            if final_sectpr is not None:
                final_sectpr.addprevious(el)
            else:
                body.append(el)

        self.mark_scope_differences()
        normalize_all_rpr(self.doc.element.body)
        normalize_all_rpr(self.doc.styles.element)
        for section in self.doc.sections:
            for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
                if not part.is_linked_to_previous:
                    normalize_all_rpr(part._element)
        self.doc.core_properties.title = self.title
        out = io.BytesIO()
        self.doc.save(out)
        for p in self.platforms:
            for w in self.imp[p].warnings:
                self.issues.append(Issue("warning", "formatting", w, p))
        return out.getvalue()

    def _retitle_heading(self, heading_el, key: str, sec: Section):
        if key == "__title__":
            set_paragraph_text(heading_el, self.title)
        elif "{product}" in key:
            text = norm_space(el_text(heading_el))
            for p in self.platforms:
                name = self.sources[p].product_name
                if name and name in text:
                    text = text.replace(name, self.product)
            set_paragraph_text(heading_el, text)

    def handle_section(self, key: str, secs: dict[Platform, Optional[Section]]) -> list:
        what = next(s.heading for s in secs.values() if s is not None) or "Document preamble"
        if any(s is not None and s.criteria_table is not None for s in secs.values()):
            return self.criteria_section(key, secs, what)
        if key.startswith("evaluation methods"):
            merged = self.evaluation_methods(secs, what)
            if merged is not None:
                return merged
        if key.startswith("scope of evaluation"):
            merged = self.scope_section(secs, what)
            if merged is None:
                merged = self.merge_simple(key, {p: (s.blocks if s else None) for p, s in secs.items()}, what)
            self._scope_tables = [b for b in merged if b.tag == qn("w:tbl")]
            return merged
        return self.merge_simple(key, {p: (s.blocks if s else None) for p, s in secs.items()}, what)

    # ---------- scope flags ----------
    def mark_scope_differences(self):
        """Highlight scope rows for pages some platforms did not evaluate, with a Word comment naming them."""
        flagged = {page_key(d["page"]): d for d in self.scope_differences}
        self.flagged_rows = 0
        if not flagged:
            return
        for tbl_el in self._scope_tables:
            table = Table(tbl_el, self.doc._body)
            for row in table.rows[1:]:
                d = flagged.get(page_key(row.cells[0].text))
                if d is None:
                    continue
                for r in row._tr.iter(qn("w:r")):
                    rpr = r.find(qn("w:rPr"))
                    if rpr is None:
                        rpr = OxmlElement("w:rPr")
                        r.insert(0, rpr)
                    for old in rpr.findall(qn("w:highlight")):
                        rpr.remove(old)
                    hl = OxmlElement("w:highlight")
                    hl.set(qn("w:val"), SCOPE_HIGHLIGHT)
                    rpr.append(hl)
                # Cell shading too: Word's comment markup paints over the text highlight of the commented cell.
                for tc in row._tr.findall(qn("w:tc")):
                    shade_cell(tc, SCOPE_SHADING)
                runs = next((p.runs for c in row.cells for p in c.paragraphs if p.runs), None)
                if runs is None:  # text only inside hyperlinks: anchor the comment on a new empty run
                    runs = [row.cells[0].paragraphs[0].add_run("")]
                self.doc.add_comment(runs, text=scope_comment(d), author=COMMENT_AUTHOR, initials="VM")
                self.flagged_rows += 1

    # ---------- evaluation methods ----------
    def evaluation_methods(self, secs, what) -> Optional[list]:
        if any(s is None for s in secs.values()):
            return None
        sigs = {p: _signature(s.blocks) for p, s in secs.items()}
        if all(sigs[p] == sigs[self.base_platform] for p in self.platforms):
            return list(secs[self.base_platform].blocks)
        # Only rewrite plain prose; anything with lists/tables keeps the labeled fallback.
        for s in secs.values():
            content = _content(s.blocks)
            if any(b.tag != qn("w:p") or is_list_para(b) for b in content):
                return None
        sentences = {}
        for p, s in secs.items():
            text = " ".join(norm_space(el_text(b)) for b in _content(s.blocks))
            sentences[p] = [x for x in re.split(r"(?<=[.!?])\s+(?=[A-Z\"'(])", text) if x]
        common = set(sentences[self.base_platform])
        for p in self.platforms[1:]:
            common &= set(sentences[p])
        base_sents = sentences[self.base_platform]
        lead = []
        for snt in base_sents:
            if snt in common:
                lead.append(snt)
            else:
                break
        trail = [snt for snt in base_sents if snt in common and snt not in lead]
        template = _content(secs[self.base_platform].blocks)[0]
        list_tmpl = self._list_template
        out = []
        if lead:
            out.append(paragraph_from_template(template, [make_run(" ".join(lead), first_run_rpr(template))],
                                               keep_numbering=False))
        for p in self.platforms:
            specific = [snt for snt in sentences[p] if snt not in common]
            tmpl = list_tmpl if list_tmpl is not None else template
            rpr = first_run_rpr(template)
            runs = self.label_runs(p, rpr) + [make_run(" ".join(specific) or MISSING_MARKER, rpr, bold=False)]
            out.append(paragraph_from_template(tmpl, runs))
        for snt in trail:
            tmpl = list_tmpl if list_tmpl is not None else template
            out.append(paragraph_from_template(tmpl, [make_run(snt, first_run_rpr(template), bold=False)]))
        out.extend(b for b in secs[self.base_platform].blocks if is_pure_sectpr_para(b))
        self.issues.append(Issue("info", "metadata-conflict",
                                 "Evaluation methods differ by platform (browser, operating system and assistive "
                                 "technology). Shared sentences are stated once; platform-specific details are listed "
                                 "under each platform's label."))
        return out

    # ---------- scope of evaluation ----------
    def scope_section(self, secs, what) -> Optional[list]:
        if any(s is None for s in secs.values()):
            return None
        parsed = {}
        for p, s in secs.items():
            info = self._split_scope(s.blocks)
            if info is None:
                return None
            parsed[p] = info
        base = parsed[self.base_platform]
        # Prose and table columns must match, otherwise fall back to per-platform labels.
        if any(_signature(parsed[p]["other"]) != _signature(base["other"])
               or parsed[p]["header_text"] != base["header_text"] for p in self.platforms):
            return None

        page_keys = {p: [k for k, _ in parsed[p]["pages"]] for p in self.platforms}
        same_pages = all(page_keys[p] == page_keys[self.base_platform] for p in self.platforms)
        comp_keys = {p: [k for k, _ in parsed[p]["components"]] for p in self.platforms}
        same_comps = all(comp_keys[p] == comp_keys[self.base_platform] for p in self.platforms)

        out = []
        for b in secs[self.base_platform].blocks:
            if b is base["table"]:
                out.append(b if same_pages else self._merged_scope_table(parsed))
            elif any(b is el for _, el in base["components"]):
                if same_comps:
                    out.append(b)
                elif b is base["components"][0][1]:
                    out.extend(self._merged_components(parsed))
            else:
                out.append(b)
        # Pages missing from a platform are flagged per page by scope.compare_scope.
        if not same_comps:
            self.issues.append(Issue("info", "scope-difference",
                                     "The tested components differ between platforms; components not tested on every "
                                     "platform are annotated with the platforms they were tested on."))
        return out

    def _split_scope(self, blocks):
        table = next((b for b in blocks if b.tag == qn("w:tbl")), None)
        if table is None:
            return None
        rows = table.findall(qn("w:tr"))
        # Web VPATs list "Page Title | URL"; native ones may list only screen names. Any width works if consistent.
        if not rows or len({len(r.findall(qn("w:tc"))) for r in rows}) != 1:
            return None
        pages = []
        for tr in rows[1:]:
            key = tuple(norm_space(el_text(c)) for c in tr.findall(qn("w:tc")))
            if any(key):
                pages.append((key, tr))
        idx = blocks.index(table)
        components = []
        for b in blocks[idx + 1:]:
            if is_list_para(b) and norm_space(el_text(b)):
                components.append((norm_space(el_text(b)), b))
            elif components and norm_space(el_text(b)):
                break
        comp_els = {id(el) for _, el in components}
        other = [b for b in blocks if b is not table and id(b) not in comp_els]
        header_text = tuple(norm_space(el_text(c)) for c in rows[0].findall(qn("w:tc")))
        return {"table": table, "header": rows[0], "header_text": header_text, "pages": pages,
                "components": components, "other": other}

    @staticmethod
    def _platform_list(platforms: list[Platform]) -> str:
        return ", ".join(p.label for p in platforms)

    def _merged_scope_table(self, parsed) -> object:
        base = parsed[self.base_platform]
        tbl = copy.deepcopy(base["table"])
        for tr in tbl.findall(qn("w:tr")):
            tbl.remove(tr)
        tbl.append(copy.deepcopy(base["header"]))

        # Every distinct row (e.g. Page Title + URL) from any platform, listed once.
        order = union_order([["\x00".join(k) for k, _ in parsed[p]["pages"]] for p in self.platforms])
        by_key = {p: {"\x00".join(k): tr for k, tr in parsed[p]["pages"]} for p in self.platforms}
        for k in sorted(order, key=lambda x: (x.split("\x00")[0].lower(), order.index(x))):
            owner = next(p for p in self.platforms if k in by_key[p])
            tbl.append(self.imp[owner](by_key[owner][k]))
        return tbl

    def _merged_components(self, parsed) -> list:
        order = union_order([[k for k, _ in parsed[p]["components"]] for p in self.platforms])
        by_key = {p: dict(parsed[p]["components"]) for p in self.platforms}
        out = []
        for k in order:
            platforms = [p for p in self.platforms if k in by_key[p]]
            el = self.imp[platforms[0]](by_key[platforms[0]][k])
            if len(platforms) < len(self.platforms):
                el.append(make_run(f" ({self._platform_list(platforms)} only)", first_run_rpr(el), bold=False))
            out.append(el)
        return out

    # ---------- conformance tables ----------
    def criteria_section(self, key, secs, what) -> list:
        mt = self.merged_tables.get(next(s.criteria_table.key for s in secs.values()
                                         if s is not None and s.criteria_table is not None))
        before = {}
        after = {}
        for p, s in secs.items():
            if s is None:
                before[p] = after[p] = None
                continue
            if s.criteria_table is None:
                before[p], after[p] = s.blocks, []
                continue
            i = s.blocks.index(s.criteria_table.tbl)
            before[p], after[p] = s.blocks[:i], s.blocks[i + 1:]
        out = self.merge_simple(key + ":before", before, f"{what} (notes before the table)")
        if mt is not None:
            out.append(self.build_table(mt))
        out.extend(self.merge_simple(key + ":after", after, f"{what} (notes after the table)"))
        return out

    def build_table(self, mt: MergedTable):
        tp = next(p for p in self.platforms if mt.tables.get(p) is not None)
        src_tbl = mt.tables[tp].tbl
        tbl = self.imp[tp](src_tbl)
        rows = tbl.findall(qn("w:tr"))
        for tr in rows[1:]:
            tbl.remove(tr)
        header = rows[0]
        trpr = header.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            header.insert(0, trpr)
        if trpr.find(qn("w:tblHeader")) is None:
            trpr.append(OxmlElement("w:tblHeader"))

        tmpl_row = src_tbl.findall(qn("w:tr"))[1] if len(src_tbl.findall(qn("w:tr"))) > 1 else header
        tmpl_cells = tmpl_row.findall(qn("w:tc"))
        for mc in mt.criteria:
            tbl.append(self._row(mc, tmpl_row, tmpl_cells))
        return tbl

    def _row(self, mc: MergedCriterion, tmpl_row, tmpl_cells):
        tr = OxmlElement("w:tr")
        trpr = tmpl_row.find(qn("w:trPr"))
        if trpr is not None:
            trpr = copy.deepcopy(trpr)
            for tag in ("w:cantSplit", "w:tblHeader"):
                for n in trpr.findall(qn(tag)):
                    trpr.remove(n)
            tr.append(trpr)
        cells = []
        for i in range(3):
            tc = OxmlElement("w:tc")
            src_pr = tmpl_cells[min(i, len(tmpl_cells) - 1)].find(qn("w:tcPr"))
            if src_pr is not None:
                tcpr = copy.deepcopy(src_pr)
                for tag in ("w:vMerge", "w:gridSpan", "w:hMerge"):
                    for n in tcpr.findall(qn(tag)):
                        tcpr.remove(n)
                tc.append(tcpr)
            tr.append(tc)
            cells.append(tc)

        owner = mc.present[0]
        for b in cell_blocks(mc.display_row.criteria_cell):
            cells[0].append(self.imp[owner](b))
        self._fill_conformance(cells[1], mc)
        self._fill_remarks(cells[2], mc)
        for tc in cells:
            blocks = cell_blocks(tc)
            if not blocks or blocks[-1].tag != qn("w:p"):
                tc.append(OxmlElement("w:p"))
        return tr

    def _fill_conformance(self, tc, mc: MergedCriterion):
        tmpl = next((p for p in cell_blocks(mc.display_row.conformance_cell) if p.tag == qn("w:p")), None)
        for i, p in enumerate(self.platforms):
            if i:
                tc.append(empty_paragraph_like(tmpl))
            row = mc.rows.get(p)
            paras = []
            if row is not None and row.conformance:
                paras = [b for b in cell_blocks(row.conformance_cell)
                         if b.tag == qn("w:p") and norm_space(el_text(b))]
            if not paras:
                tc.append(self.label_paragraph(p, tmpl, MISSING_MARKER, italic=True))
                continue
            imported = [self.imp[p](b) for b in paras]
            insert_runs_at_start(imported[0], self.label_runs(p, first_run_rpr(imported[0])))
            for b in imported:
                tc.append(b)

    def _fill_remarks(self, tc, mc: MergedCriterion):
        tmpl = next((p for p in cell_blocks(mc.display_row.remarks_cell)
                     if p.tag == qn("w:p") and not is_list_para(p)), None)
        for i, p in enumerate(self.platforms):
            if i:
                tc.append(empty_paragraph_like(tmpl))
            row = mc.rows.get(p)
            if row is None:
                tc.append(self.label_paragraph(p, tmpl, MISSING_MARKER, italic=True))
                continue
            blocks = [self.imp[p](b) for b in _content(cell_blocks(row.remarks_cell))]
            local_tmpl = next((b for b in cell_blocks(row.remarks_cell) if b.tag == qn("w:p")), tmpl)
            for b in self.labeled_blocks(p, blocks, local_tmpl if not is_list_para(local_tmpl) else tmpl):
                tc.append(b)
