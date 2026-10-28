"""Parse a single-platform VPAT .docx into sections and criteria rows."""
from __future__ import annotations

import io
import re
import zipfile
from typing import Optional

from docx import Document
from docx.oxml.ns import qn

from .docx_utils import block_texts, cell_blocks, el_text, norm_space, para_style_id
from .models import (
    VALID_CONFORMANCE,
    CriteriaTable,
    CriterionRow,
    Issue,
    Platform,
    Section,
    SourceVpat,
)

SC_RE = re.compile(r"^\s*(\d+\.\d+\.\d+)\s+(.+?)\s*(?:\((Level\s+[^)]*)\))?\s*$", re.IGNORECASE)
TITLE_SUFFIX_RE = re.compile(r"\s*Accessibility Conformance Report\s*$", re.IGNORECASE)


class VpatParseError(Exception):
    pass


def load_document(data: bytes, filename: str):
    if not filename.lower().endswith(".docx"):
        raise VpatParseError(f"{filename} is not a .docx file.")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise VpatParseError(f"{filename} is not a valid Word document (it is not a .docx package).")
    try:
        return Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of errors for corrupt packages
        raise VpatParseError(f"{filename} could not be opened as a Word document: {exc}") from exc


def _heading_levels(doc) -> dict[str, int]:
    """Map paragraph style ids to heading levels (works with localized style ids)."""
    levels = {}
    for style in doc.styles.element.findall(qn("w:style")):
        sid = style.get(qn("w:styleId"))
        name_el = style.find(qn("w:name"))
        name = (name_el.get(qn("w:val")) if name_el is not None else sid or "").lower()
        m = re.match(r"heading\s*(\d)", name) or re.match(r"heading(\d)", (sid or "").lower())
        if m:
            levels[sid] = int(m.group(1))
    return levels


def section_key(heading: str, product_name: str) -> str:
    text = norm_space(heading)
    if product_name:
        text = text.replace(product_name, "{product}")
    return text.rstrip(":").strip().lower()


def is_criteria_header(tbl) -> bool:
    rows = tbl.findall(qn("w:tr"))
    if not rows:
        return False
    cells = rows[0].findall(qn("w:tc"))
    texts = [norm_space(el_text(c)).lower() for c in cells]
    return (
        len(texts) >= 3
        and "criteria" in texts[0]
        and "conformance" in texts[1]
        and "remarks" in texts[2]
    )


def normalize_conformance(value: str) -> Optional[str]:
    """Return the canonical conformance term if `value` is one, else None."""
    v = norm_space(value).lower()
    for term in VALID_CONFORMANCE:
        if v == term.lower():
            return term
    return None


def parse_criteria_table(tbl, heading: str, report_heading: Optional[str], src: SourceVpat) -> CriteriaTable:
    table = CriteriaTable(key=section_key(heading, src.product_name), heading=heading,
                          report_heading=report_heading, tbl=tbl)
    seen: dict[str, CriterionRow] = {}
    for tr in tbl.findall(qn("w:tr"))[1:]:
        cells = tr.findall(qn("w:tc"))
        if not any(norm_space(el_text(c)) for c in cells):
            continue
        if len(cells) < 3:
            src.issues.append(Issue("error", "structure",
                                    f"A row in '{heading}' has {len(cells)} cells instead of 3 and was skipped: "
                                    f"'{norm_space(el_text(tr))[:80]}'", src.platform))
            continue
        title = norm_space(el_text(cells[0]))
        m = SC_RE.match(title)
        if m:
            sc_id, level = m.group(1), m.group(3)
            key = sc_id
            malformed = False
        else:
            sc_id, level = None, None
            key = "text:" + title.lower()
            malformed = True
            src.issues.append(Issue("warning", "malformed-criterion",
                                    f"Criterion '{title[:80]}' has no recognizable WCAG identifier; it was matched by "
                                    f"its exact text only.", src.platform, title[:40]))
        conformance = norm_space(el_text(cells[1]))
        row = CriterionRow(
            key=key, sc_id=sc_id, title=title, level=level, conformance=conformance, tr=tr,
            criteria_cell=cells[0], conformance_cell=cells[1], remarks_cell=cells[2],
            remarks_text=block_texts(cell_blocks(cells[2])), malformed=malformed,
        )
        if key in seen:
            table.duplicates.append(row)
            src.issues.append(Issue("error", "duplicate-criterion",
                                    f"{title} appears more than once in '{heading}'. Only the first row is merged; "
                                    f"the duplicate says '{conformance or '(blank)'}'. Fix the source and regenerate.",
                                    src.platform, key))
            continue
        seen[key] = row
        table.rows.append(row)
    return table


def parse_vpat(data: bytes, filename: str, platform: Platform) -> SourceVpat:
    doc = load_document(data, filename)
    levels = _heading_levels(doc)
    body = doc.element.body

    sections: list[Section] = []
    current = Section(key="__preamble__", heading="", level=0, heading_el=None)
    sections.append(current)
    for el in body.iterchildren():
        if el.tag == qn("w:sectPr"):
            continue
        if el.tag == qn("w:p"):
            lvl = levels.get(para_style_id(el))
            text = norm_space(el_text(el))
            if lvl and text:
                current = Section(key="", heading=text, level=lvl, heading_el=el)
                sections.append(current)
                continue
        current.blocks.append(el)
    if not sections[0].blocks:
        sections.pop(0)

    title_section = next((s for s in sections if s.level == 1), None)
    product = ""
    if title_section is not None:
        product = TITLE_SUFFIX_RE.sub("", title_section.heading).strip()
    if not product:
        name_sec = next((s for s in sections if s.heading.lower().startswith("name of product")), None)
        if name_sec:
            texts = block_texts(name_sec.blocks)
            product = texts[0] if texts else ""

    src = SourceVpat(platform=platform, filename=filename, document=doc, product_name=product,
                     sections=sections, tables=[])

    report_heading = None
    for s in sections:
        s.key = "__title__" if s is title_section else section_key(s.heading, product)
        if s.level == 2:
            report_heading = s.heading
        for b in s.blocks:
            if b.tag == qn("w:tbl") and is_criteria_header(b):
                if s.criteria_table is not None:
                    src.issues.append(Issue("warning", "structure",
                                            f"Section '{s.heading}' holds more than one conformance table; "
                                            f"only the first is merged.", platform))
                    continue
                table = parse_criteria_table(b, s.heading or "Conformance table", report_heading, src)
                s.criteria_table = table
                src.tables.append(table)

    # Disambiguate section keys that repeat (e.g. two "Notes:" headings).
    counts: dict[str, int] = {}
    for s in sections:
        n = counts.get(s.key, 0)
        counts[s.key] = n + 1
        if n:
            s.key = f"{s.key}#{n + 1}"

    if not src.tables:
        raise VpatParseError(
            f"{filename} does not contain a VPAT conformance table "
            f"(a table with 'Criteria', 'Conformance Level' and 'Remarks and Explanations' columns)."
        )
    return src


PLATFORM_HINTS = {
    Platform.IOS: [r"\bios\b", r"safari", r"voice\s?over", r"mwebios", r"iphone", r"ipad"],
    Platform.ANDROID: [r"android", r"talk\s?back"],
    Platform.DESKTOP: [r"windows", r"\bnvda\b", r"\bjaws\b", r"macos", r"desktop"],
}


NATIVE_HINTS = [r"\bnative\b", r"\bapps?\b", r"\.ipa\b", r"\.apk\b", r"xcode", r"android studio",
                r"accessibility (scanner|inspector)"]
WEB_HINTS = [r"mweb", r"mobile[\s_-]?web", r"\bweb\b", r"browser", r"safari", r"chrome", r"firefox", r"\bedge\b"]


def _matches(patterns, text) -> bool:
    return any(re.search(pat, text) for pat in patterns)


def guess_platform(src: SourceVpat) -> Optional[Platform]:
    """Best-effort guess of OS and native-vs-web from the evaluation methods and file/product name."""
    eval_sec = next((s for s in src.sections if s.key.startswith("evaluation methods")), None)
    eval_text = " ".join(block_texts(eval_sec.blocks)).lower() if eval_sec else ""
    name_text = re.sub(r"[_\-]+", " ", f"{src.filename} {src.product_name}".lower())

    os_guess = None
    for text in (eval_text, name_text):
        # VoiceOver/Safari also exist on macOS; a Mac-only document is desktop.
        if re.search(r"mac\s?os|macbook", text) and not re.search(r"\bios\b|iphone|ipad", text):
            os_guess = Platform.DESKTOP
            break
        hits = [p for p, pats in PLATFORM_HINTS.items() if any(re.search(pat, text) for pat in pats)]
        # Mobile hints win over desktop ones ("Chrome" etc. appear everywhere).
        os_guess = next((p for p in (Platform.IOS, Platform.ANDROID, Platform.DESKTOP) if p in hits), None)
        if os_guess:
            break
    if os_guess in (None, Platform.DESKTOP):
        return os_guess

    # The file/product name is the most explicit signal ("…-Native Mobile iOS", "…-MWEBIOS-…").
    if _matches(NATIVE_HINTS, name_text):
        native = True
    elif _matches(WEB_HINTS, name_text):
        native = False
    else:
        native = _matches(NATIVE_HINTS, eval_text) and not _matches(WEB_HINTS, eval_text)
    if native:
        return Platform.IOS_NATIVE if os_guess is Platform.IOS else Platform.ANDROID_NATIVE
    return os_guess
